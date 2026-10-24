"""
Human batch 2 (guide's feedback of 15 Oct). No GPU. Same annotation sheet and
the same rules as Experiment 2; two annotators label every item, blind.

What goes in (shuffled together, so nobody can tell which group an item is from):
  faithful  30 per model   correction 6: are the cases scored "faithful" really
                           faithful? Random, stratified by source (NQ/TriviaQA/PopQA)
  f3        30 per model   extra F3 labels: every unlabelled apparent failure the
                           rebuilt classifier calls F3, topped up by the judge's
                           probability that the passage lacks the answer
  pilot     all            enrichment-pilot F1 candidates (gold and mixed passages)
  probe     5 per model x  other wrong pilot answers (partial / F6 / F2 / F3): how
            condition      many real overrides the pilot classifier missed

Steps:
  python Tests/experiment2b_batch2.py sample      -> sheets for the two annotators
  python Tests/experiment2b_batch2.py agreement   -> after both sheets are back
  python Tests/experiment2b_batch2.py report      -> after adjudication.csv is filled

Files: data/final/Phase_05/exp2b/   (HIDDEN_key.csv must not reach the annotators --
the repo is public, so do not push it until both sheets are back)
"""

import argparse
import csv
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path

# ---- copied unchanged from experiment2_annotation.py, so batch 2 uses exactly the batch-1
# ---- sheet, allowed values, normalisations and adjudication rules
TYPES = ["F1", "F2", "F3", "F4", "F5", "F6"]
MAX_PASSAGE_CHARS = 30000           # Excel cell limit is 32,767
ANNOT_COLS = ["failure_type", "passage_contains_answer", "answer_correct",
              "answer_supported_by_passage", "answer_conflicts_with_passage",
              "confidence", "rationale"]
ALLOWED = {
    "failure_type": set(TYPES),
    "passage_contains_answer": {"yes", "no", "unclear"},
    "answer_correct": {"yes", "no", "partially", "unclear"},
    "answer_supported_by_passage": {"yes", "no", "partially", "unclear"},
    "answer_conflicts_with_passage": {"yes", "no", "unclear"},
    "confidence": {"high", "medium", "low"},
}

GUIDELINES = """ANNOTATION GUIDELINES -- Retrieval failure types (read fully before starting)

You see: question, gold answer(s), retrieved passage(s), the model's answer.
Work ALONE. Do not discuss cases with the other annotator until both sheets are done.
Fill every annotation column. Allowed values are listed below; use them exactly.

Answer these in order, then pick ONE failure_type:
  passage_contains_answer       yes / no / unclear
      Does the passage explicitly state the gold answer (any valid form)?
  answer_correct                yes / no / partially / unclear
      Is the model's answer factually correct for the question?
  answer_supported_by_passage   yes / no / partially / unclear
      Is what the model said stated in the passage (even if it's the wrong fact)?
  answer_conflicts_with_passage yes / no / unclear
      Does the model's answer contradict what the passage says about the question?

failure_type (exactly one):
  F1  Genuine parametric override -- passage contains the correct answer, the model's
      answer conflicts with it, and the answer is NOT taken from the passage, i.e. it is
      better explained by the model's own prior knowledge.
  F2  Wrong extraction / wrong interpretation -- the model uses the passage but takes the
      wrong fact, answers a different aspect of the question, or misreads the evidence.
  F3  Evidence absent -- the passage does not actually support the gold answer.
  F4  Partially correct / incomplete -- part of the correct answer is present but
      something required is missing or altered.
  F5  Actually correct -- the answer is right; the automatic scoring was too strict.
  F6  Ambiguous -- you cannot reliably decide among the above.

CRITICAL: a wrong answer is NOT automatically F1. Choose F1 only when the passage has the
answer AND the model's answer is not drawn from the passage. If the model's wrong answer
appears in the passage, that is F2. If you can't tell, use F6.

confidence: high / medium / low.   rationale: one short sentence.
"""


def read_csv(p):
    with open(p, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(p, rows, cols):
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def allocate(cells, target):
    """Equal share per cell, capped by availability; leftovers redistributed."""
    quota = {k: 0 for k in cells}
    remaining = target
    open_cells = [k for k in cells if cells[k] > 0]
    while remaining > 0 and open_cells:
        share = max(1, remaining // len(open_cells))
        for k in list(open_cells):
            give = min(share, cells[k] - quota[k], remaining)
            quota[k] += give
            remaining -= give
            if quota[k] >= cells[k]:
                open_cells.remove(k)
            if remaining == 0:
                break
    return quota


def kappa(x, y, labels):
    n = len(x)
    if n == 0:
        return float("nan")
    po = sum(a == b for a, b in zip(x, y)) / n
    cx, cy = Counter(x), Counter(y)
    pe = sum(cx[l] * cy[l] for l in labels) / (n * n)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def boot_kappa(x, y, labels, reps=2000):
    rng = random.Random(0)
    n, vals = len(x), []
    for _ in range(reps):
        idx = [rng.randrange(n) for _ in range(n)]
        k = kappa([x[i] for i in idx], [y[i] for i in idx], labels)
        if not math.isnan(k):
            vals.append(k)
    vals.sort()
    return vals[int(0.025 * len(vals))], vals[int(0.975 * len(vals)) - 1]


def matrix(x, y, labels):
    m = {a: Counter() for a in labels}
    for a, b in zip(x, y):
        m[a][b] += 1
    return m


def print_matrix(m, labels, rlab, clab):
    corner = f"{rlab} / {clab}"
    print(f"{corner:<14}" + "".join(f"{l:>6}" for l in labels))
    for a in labels:
        print(f"{a:<14}" + "".join(f"{m[a][b]:>6}" for b in labels))


# Fixed normalizations for values outside the allowed set. Applied by the script
# (the annotators' files are never edited) and printed, so they can be reported.
NORMALIZE = {
    "passage_contains_answer": {"partially": "unclear", "partial": "unclear"},
    "answer_conflicts_with_passage": {"partially": "unclear", "partial": "unclear"},
    "answer_correct": {"partialy": "partially", "partial": "partially"},
    "answer_supported_by_passage": {"partialy": "partially", "partial": "partially"},
}

ADJUDICATION_RULES = """ADJUDICATION RULES -- fixed BEFORE any disagreement is resolved.
Adjudicator: a third person (not annotator A or B) with the passage open.
Apply the checks in this order; the first that fits decides the final type.

  1. F5  The answer is correct for the question: the same entity or fact in different
         words or an alias ("Jack Lord" = John Joseph Patrick Ryan, "Arthur" = King Arthur),
         including when the gold answer itself is demonstrably wrong.
  2. F4  The answer is right but less specific than required: only the year of a full
         date ("1967" for "July 1, 1967"), a broader category ("Football" for "American
         football"), or only part of a multi-part answer.
  3. F3  The passage does not state the gold answer (or any valid answer).
  4. F6  The model gives no answer: repeats the question, placeholder text, or a
         truncated fragment with no answer in it.
  5. F2  The model's wrong answer, or the fact it states, is taken from the passage --
         including copying a general description that does not answer the question.
  6. F1  The passage states the gold answer AND the model's wrong answer does not appear
         in the passage.
  7. F6  Still undecidable.

If the adjudicator cannot decide with these rules, the guide decides.
Write the result in final_type and 'adjudicated' in how. Do not change agreed rows.
"""


def validate_sheet(rows, name):
    """Applies the fixed normalizations (reported), returns the ids of rows the
    annotator left without a failure_type, and stops on anything else invalid."""
    bad, fixes, unannotated = [], Counter(), []
    for r in rows:
        if not (r.get("failure_type") or "").strip():
            r["failure_type"] = ""
            unannotated.append(r["item_id"])
            continue
        for c, ok in ALLOWED.items():
            v = (r.get(c) or "").strip()
            v = v.upper() if c == "failure_type" else v.lower()
            if v in NORMALIZE.get(c, {}):
                fixes[(c, v, NORMALIZE[c][v])] += 1
                v = NORMALIZE[c][v]
            r[c] = v
            if v not in ok:
                bad.append((r["item_id"], c, v))
    for (c, v, nv), n in sorted(fixes.items()):
        print(f"Sheet {name}: normalized {n} x {c} '{v}' -> '{nv}' (fixed rule, report it)")
    if unannotated:
        print(f"Sheet {name}: {len(unannotated)} rows not annotated: {', '.join(sorted(unannotated))}")
    if bad:
        print(f"Sheet {name}: {len(bad)} invalid entries, e.g. {bad[:5]}")
        raise SystemExit("Fix these before computing agreement.")
    return unannotated


SEED = 2026
MODELS = ("gemma", "llama")
CONDS = ("gold", "mixed")
SEP = "\n\n---\n\n"
SHEET_COLS = ["item_id", "question", "gold_answers", "model_answer", "retrieved_passages"] + ANNOT_COLS
PROBE_LABELS = ("partial", "F6", "F2", "F3")
NOTE = """BATCH 2 NOTE -- the rules below are exactly the batch-1 rules. In addition:
  * This batch mixes different kinds of cases, and some model answers are correct.
    If the model's answer is right, choose F5.
  * Some items show one passage, others several separated by ---. Judge against all of them.
  * The same question can appear more than once (different model or different passages).
    Judge every item on its own.

"""


# ------------------------------------------------------------------ helpers
def load_jsonl(p):
    p = Path(p)
    return [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()] if p.exists() else []


def ptext(passages):
    t = SEP.join(passages)
    return t if len(t) <= MAX_PASSAGE_CHARS else t[:MAX_PASSAGE_CHARS] + " [...TRUNCATED]"


def proportional(sizes, n):
    """Largest-remainder allocation of n proportional to sizes."""
    tot = sum(sizes.values())
    if tot <= n:
        return dict(sizes)
    raw = {k: n * s / tot for k, s in sizes.items()}
    q = {k: int(v) for k, v in raw.items()}
    for k in sorted(raw, key=lambda k: (-(raw[k] - q[k]), k))[: n - sum(q.values())]:
        q[k] += 1
    return q


def classify(J):                      # identical to experiment1b_rebuilt_classifier.classify
    c = J.get("correct")
    if c is None:
        return None
    if c == "yes":
        return "F5"
    if c == "partial":
        return "F4"
    if J.get("contains") == "no":
        return "F3"
    if J.get("answers") == "no":
        return "F6"
    if J.get("supported") == "yes":
        return "F2"
    if J.get("same") == "yes" and J.get("noctx_correct") != "yes":
        return "F1"
    return "F6"


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


def dataset_of(source):
    return "popqa" if source == "popqa" else "natural_rag_other"


def pilot_labels(pilot_dir):
    """Recompute the pilot classification exactly as experiment1b --task pilot does."""
    d = Path(pilot_dir)
    pool = {r["pid"]: r for r in load_jsonl(d / "pilot_pool.jsonl")}
    jl = {tuple(r["key"]): r["label"] for r in load_jsonl(d / "judge_pilot.jsonl")}
    labels, wps, cbs, n_pool = {}, {}, {}, {}
    for m in MODELS:
        cb = load_jsonl(d / f"pilot_closedbook_{m}.jsonl")
        wp = {(r["pid"], r["cond"]): r for r in load_jsonl(d / f"pilot_withpassage_{m}.jsonl")}
        n_pool[m] = len(cb)
        wrong = [r for r in cb if jl.get((r["pid"], m, "-", "cb_correct")) == "no" and not r["cb_string_correct"]]
        for r in wrong:
            cbs[(r["pid"], m)] = r
            for cond in CONDS:
                w = wp.get((r["pid"], cond))
                if not w:
                    continue
                c = "yes" if w["wp_string_correct"] else jl.get((r["pid"], m, cond, "wp_correct"))
                J = {"correct": c, "noctx_correct": "no",
                     "contains": jl.get((r["pid"], "-", cond, "contains")),
                     "answers": jl.get((r["pid"], m, cond, "answers")),
                     "supported": jl.get((r["pid"], m, cond, "supported")),
                     "same": jl.get((r["pid"], m, cond, "same"))}
                t = classify(J)
                if t:
                    labels[(r["pid"], m, cond)] = {"F5": "fixed_by_passage", "F4": "partial"}.get(t, t)
                    wps[(r["pid"], m, cond)] = w
    return pool, labels, wps, cbs, n_pool


def passages_for(p, cond):
    if cond == "gold":
        return [p["gold_passage"]]
    ps = list(p["distractors"])
    ps.insert(p["gold_position"], p["gold_passage"])
    return ps


# ------------------------------------------------------------------ sheets (xlsx with drop-downs, csv fallback)
def write_sheet(path_noext, rows, guidelines):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        with open(f"{path_noext}.csv", "w", newline="", encoding="utf-8-sig") as f:    # BOM: Excel reads UTF-8
            w = csv.DictWriter(f, fieldnames=SHEET_COLS, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        return f"{path_noext}.csv"
    wb = Workbook()
    ws = wb.active
    ws.title = "annotate"
    ws.append(SHEET_COLS)
    for r in rows:
        ws.append([r.get(c, "") for c in SHEET_COLS])
    widths = {"item_id": 9, "question": 40, "gold_answers": 28, "model_answer": 30, "retrieved_passages": 90,
              "rationale": 45}
    for i, c in enumerate(SHEET_COLS, 1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = widths.get(c, 16)
        ws.cell(row=1, column=i).font = Font(bold=True)
        if c in ALLOWED:
            dv = DataValidation(type="list", formula1='"' + ",".join(sorted(ALLOWED[c])) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Not allowed",
                                error="Pick one of the listed values (see the guidelines tab).")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(rows) + 1}")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "B2"
    g = wb.create_sheet("guidelines")
    g.column_dimensions["A"].width = 110
    for line in guidelines.splitlines():
        g.append([line])
    wb.save(f"{path_noext}.xlsx")
    return f"{path_noext}.xlsx"


def read_table(p):
    """Rows of a returned sheet. The format is taken from the file's content, not its name
    (an Excel file saved as '.csv' is still read correctly)."""
    with open(p, "rb") as f:
        is_xlsx = f.read(4) == b"PK\x03\x04"
    if not is_xlsx:
        return read_csv(p)
    import io
    from openpyxl import load_workbook
    with open(p, "rb") as f:
        ws = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    return [{h: ("" if v is None else str(v)) for h, v in zip(head, vals)} for vals in it
            if vals and any(v is not None for v in vals)]


def read_sheet(d, name):
    found = [p for p in (Path(d) / f"annotation_sheet_{name}.xlsx", Path(d) / f"annotation_sheet_{name}.csv")
             if p.exists()]
    if not found:
        raise SystemExit(f"annotation_sheet_{name}.xlsx / .csv not found in {d}")
    p = max(found, key=lambda q: q.stat().st_mtime)          # the returned copy is the newest file
    rows = read_table(p)
    filled = sum(bool((r.get("failure_type") or "").strip()) for r in rows)
    if filled < len(rows) / 2:
        raise SystemExit(f"{p.name}: only {filled} of {len(rows)} rows have a failure_type -- this looks like the "
                         "blank sheet. Copy the annotator's returned file over it and run again.")
    return rows, p


def norm(t):
    t = "".join(ch for ch in (t or "").lower() if ch not in set('!"#$%&\'()*+,-./:;<=>?@[\\]^_`{|}~'))
    t = " ".join(w for w in t.split() if w not in ("a", "an", "the"))
    return t


def found_in(text, hay):
    t = norm(text)
    return "yes" if t and f" {t} " in f" {norm(hay)} " else "no"


def write_adjudication_xlsx(path, rows):
    """Adjudication workbook: disagreements first (grouped by the label pair), agreed rows below.
    final_type has a drop-down; the two helper columns are plain string matches."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    cols = ["item_id", "how", "annot_A", "annot_B", "final_type", "adjudication_note", "model_answer_found_in_passage",
            "a_gold_answer_found_in_passage", "A_contains", "B_contains", "A_correct", "B_correct", "A_supported",
            "B_supported", "A_conflicts", "B_conflicts", "rationale_A", "rationale_B", "question", "gold_answers",
            "model_answer", "retrieved_passages"]
    if any(r.get("claude_suggestion") for r in rows):          # optional AI-drafted suggestions, kept separate
        k = cols.index("final_type") + 1
        cols = cols[:k] + ["claude_suggestion", "claude_reason"] + cols[k:]
    order = sorted(rows, key=lambda r: (r["how"] == "agreed", r["annot_A"], r["annot_B"], r["item_id"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "adjudicate"
    ws.append(cols)
    for r in order:
        ws.append([r.get(c, "") for c in cols])
    widths = {"item_id": 9, "how": 22, "final_type": 11, "claude_suggestion": 11, "claude_reason": 50,
              "adjudication_note": 40, "rationale_A": 45, "rationale_B": 45,
              "question": 40, "gold_answers": 25, "model_answer": 30, "retrieved_passages": 90}
    todo = PatternFill("solid", fgColor="FFF2CC")
    for i, c in enumerate(cols, 1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = widths.get(c, 12)
        ws.cell(row=1, column=i).font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    fcol = ws.cell(row=1, column=cols.index("final_type") + 1).column_letter
    dv.add(f"{fcol}2:{fcol}{len(order) + 1}")
    wrap = Alignment(wrap_text=True, vertical="top")
    for k, row in enumerate(ws.iter_rows(min_row=2, max_row=len(order) + 1)):
        for cell in row:
            cell.alignment = wrap
        if order[k]["how"] != "agreed":
            row[cols.index("final_type")].fill = todo
    ws.freeze_panes = "F2"
    g = wb.create_sheet("rules")
    g.column_dimensions["A"].width = 110
    for line in (ADJUDICATION_RULES + "\nHOW TO USE THIS SHEET\n"
                 "  Fill final_type (yellow cells) for every row that is not 'agreed'; add a short note.\n"
                 "  Do not change agreed rows. Save the file under the same name.\n"
                 "  The two helper columns are exact string matches after lower-casing and removing punctuation:\n"
                 "  read the passage before deciding -- a paraphrase or a long answer is rarely an exact match.\n"
                 "  Do not open HIDDEN_key.csv before adjudication is finished.\n"
                 + ("  claude_suggestion / claude_reason: rule-based suggestions drafted by Claude (an AI assistant).\n"
                    "  They are NOT labels. Decide each row yourself; the report counts how often you agreed.\n"
                    if "claude_suggestion" in cols else "")).splitlines():
        g.append([line])
    wb.save(path)


# ------------------------------------------------------------------ 1) sample
def cmd_sample(a):
    out = Path(a.outdir)
    if any(out.glob("annotation_sheet_*")):
        raise SystemExit(f"{out} already has annotation sheets -- refusing to overwrite them. "
                         "Delete the folder only if the annotators have not started.")
    rng = random.Random(SEED)
    ret = {r["query_id"]: r.get("retrieved_passages") or [] for r in load_jsonl(a.retrieval)}
    exp1 = {(r["question_id"], r["model"]): r for r in load_jsonl(a.exp1)}
    labelled = {(r["question_id"], r["model"]) for r in load_jsonl(a.validated)}
    items = []

    # faithful: both labels faithful, not an apparent failure, ConFiQA excluded (as in Experiment 1)
    data = load_jsonl(a.dataset)
    for m in MODELS:
        pop = defaultdict(list)
        for x in data:
            if x["source"] != "confiqa" and (x["query_id"], m) not in exp1 \
                    and x.get(f"{m}_label_full") == "faithful" and x.get(f"{m}_label") == "faithful":
                pop[x["source"]].append(x)
        sizes = {k: len(v) for k, v in pop.items()}
        quota = proportional(sizes, a.n_faithful)
        for src in sorted(pop):
            cand = sorted(pop[src], key=lambda x: x["query_id"])
            rng.shuffle(cand)
            for x in cand[:quota[src]]:
                items.append({"group": "faithful", "question_id": x["query_id"], "model": m, "cond": "",
                              "source": src, "dataset": dataset_of(src), "selection": "random_stratified",
                              "auto_type": "faithful", "sampling_weight": round(sizes[src] / quota[src], 4),
                              "question": x["question"], "gold_answers": x["gold_answer_text"],
                              "model_answer": x[f"{m}_answer"], "retrieved_passages": ptext(ret.get(x["query_id"], []))})

    # f3: unlabelled apparent failures predicted F3, topped up by P(passage lacks the answer)
    preds = {(r["question_id"], r["model"]): r for r in read_csv(a.predictions)}
    p_no = {r["key"][0]: float(r["probs"].get("no", 0.0)) for r in load_jsonl(a.judge_existing)
            if r.get("kind") == "contains"}
    for m in MODELS:
        unl = sorted(k for k in exp1 if k[1] == m and k not in labelled and k in preds)
        f3 = [k for k in unl if preds[k]["rebuilt_type"] == "F3"]
        rest = sorted((k for k in unl if preds[k]["rebuilt_type"] != "F3"), key=lambda k: (-p_no.get(k[0], 0.0), k))
        if len(f3) > a.n_f3:
            rng.shuffle(f3)
            chosen = [(k, "predicted_F3") for k in f3[:a.n_f3]]
        else:
            chosen = [(k, "predicted_F3") for k in f3] + [(k, "top_up_p_contains_no") for k in rest[:a.n_f3 - len(f3)]]
        for k, why in chosen:
            r = exp1[k]
            items.append({"group": "f3", "question_id": k[0], "model": m, "cond": "",
                          "source": k[0].split("_")[0], "dataset": r["dataset"], "selection": why,
                          "auto_type": r["failure_type"], "rebuilt_type": preds[k]["rebuilt_type"],
                          "p_contains_no": round(p_no.get(k[0], float("nan")), 4), "noctx_answer": r.get("noctx_answer", ""),
                          "question": r["question"], "gold_answers": r["gold"], "model_answer": r["model_answer"],
                          "retrieved_passages": ptext(ret.get(k[0], []))})

    # pilot: every F1 candidate, plus a probe of other wrong answers (where missed overrides would hide)
    pool, labels, wps, cbs, n_pool = pilot_labels(a.pilot_dir)
    listed = {(r["pid"], r["model"], r["cond"]) for r in read_csv(Path(a.pilot_dir) / "pilot_F1_candidates.csv")}
    cands = sorted(k for k, t in labels.items() if t == "F1")
    if set(cands) != listed:
        raise SystemExit(f"recomputed pilot candidates ({len(cands)}) differ from pilot_F1_candidates.csv "
                         f"({len(listed)}) -- the pilot files are inconsistent")

    def pilot_item(k, group, selection, weight, label):
        pid, m, cond = k
        p = pool[pid]
        return {"group": group, "question_id": pid, "model": m, "cond": cond, "source": p["source"],
                "dataset": "pilot_" + p["source"], "selection": selection, "auto_type": label,
                "sampling_weight": weight, "closed_book_answer": cbs[(pid, m)]["cb_answer"],
                "cb_mean_logprob": cbs[(pid, m)]["cb_mean_logprob"],
                "question": p["question"], "gold_answers": " | ".join(p["answers"]),
                "model_answer": wps[k]["wp_answer"], "retrieved_passages": ptext(passages_for(p, cond))}

    for k in cands:
        items.append(pilot_item(k, "pilot", "all_F1_candidates", 1.0, "F1"))
    for m in MODELS:
        for cond in CONDS:
            strata = {lab: sorted(k for k, t in labels.items() if k[1] == m and k[2] == cond and t == lab)
                      for lab in PROBE_LABELS}
            quota = allocate({lab: len(v) for lab, v in strata.items()}, a.n_probe)
            for lab in PROBE_LABELS:
                v = strata[lab]
                rng.shuffle(v)
                for k in v[:quota[lab]]:
                    items.append(pilot_item(k, "probe", f"probe_{lab}", round(len(v) / quota[lab], 4), lab))

    # blind: one shuffle for the ids, a different order for each annotator
    rng.shuffle(items)
    for i, it in enumerate(items, 1):
        it["item_id"] = f"b2_{i:03d}"
        for c in ANNOT_COLS:
            it[c] = ""
    out.mkdir(parents=True, exist_ok=True)
    guide = NOTE + GUIDELINES
    files = []
    for name, s in (("A", SEED + 11), ("B", SEED + 12)):
        sheet = items[:]
        random.Random(s).shuffle(sheet)
        files.append(write_sheet(str(out / f"annotation_sheet_{name}"), sheet, guide))
    key_cols = ["item_id", "group", "question_id", "model", "cond", "source", "dataset", "selection", "auto_type",
                "rebuilt_type", "p_contains_no", "sampling_weight", "noctx_answer", "closed_book_answer",
                "cb_mean_logprob"]
    write_csv(out / "HIDDEN_key.csv", [{c: it.get(c, "") for c in key_cols} for it in items], key_cols)
    (out / "annotation_guidelines.txt").write_text(guide, encoding="utf-8")
    (out / ".gitignore").write_text("HIDDEN_key.csv\n", encoding="utf-8")
    json.dump({"seed": SEED, "n_pool_pilot": n_pool, "n_items": len(items)}, open(out / "sample_info.json", "w"), indent=2)

    print(f"Batch 2: {len(items)} items")
    c = Counter((it["group"], it["model"]) for it in items)
    for g in ("faithful", "f3", "pilot", "probe"):
        print(f"  {g:<9} gemma {c[(g, 'gemma')]:>3}   llama {c[(g, 'llama')]:>3}")
    print("  f3 selection:", dict(Counter((it["model"], it["selection"]) for it in items if it["group"] == "f3")))
    print("  pilot by condition:", dict(Counter((it["model"], it["cond"]) for it in items if it["group"] == "pilot")))
    print("  faithful by source:", dict(Counter((it["model"], it["source"]) for it in items if it["group"] == "faithful")))
    tab = " (the sheets also have a guidelines tab)" if files[0].endswith(".xlsx") else ""
    print(f"\nSend {Path(files[0]).name} to annotator A and {Path(files[1]).name} to annotator B, with "
          f"annotation_guidelines.txt{tab}.")
    print("Do NOT share or push HIDDEN_key.csv until both sheets are back (the repo is public).")


# ------------------------------------------------------------------ 2) agreement
def cmd_agreement(a):
    d = Path(a.outdir)
    A_rows, pa = read_sheet(d, "A")
    A = {r["item_id"]: r for r in A_rows}
    validate_sheet(list(A.values()), "A")
    if a.single:
        final = [{"item_id": i, "annot_A": A[i]["failure_type"], "annot_B": "",
                  "final_type": A[i]["failure_type"], "how": "single_annotator"} for i in sorted(A)]
        write_csv(d / "adjudication.csv", final, list(final[0].keys()))
        print("Single annotator: sheet A taken as final. Report this as a limitation.")
        return
    B_rows, pb = read_sheet(d, "B")
    B = {r["item_id"]: r for r in B_rows}
    validate_sheet(list(B.values()), "B")
    print(f"read {pa.name} and {pb.name}")
    common = sorted(set(A) & set(B))
    edited = [i for i in common if any((A[i].get(c) or "").strip() != (B[i].get(c) or "").strip()
                                       for c in ("question", "gold_answers", "model_answer"))]
    if edited:
        print(f"note: {len(edited)} items where the two sheets' question/gold/answer text differs (a cell was edited or "
              f"auto-formatted by Excel): {', '.join(edited)} -- the adjudication sheet shows sheet B's text")
    ids = [i for i in common if A[i]["failure_type"] and B[i]["failure_type"]]
    partial = [i for i in common if i not in ids]
    if not ids:
        raise SystemExit("no item was annotated by both annotators")
    x, y = [A[i]["failure_type"] for i in ids], [B[i]["failure_type"] for i in ids]
    k = kappa(x, y, TYPES)
    lo, hi = boot_kappa(x, y, TYPES)
    xb, yb = ["F1" if v == "F1" else "other" for v in x], ["F1" if v == "F1" else "other" for v in y]
    kb = kappa(xb, yb, ["F1", "other"])
    lob, hib = boot_kappa(xb, yb, ["F1", "other"])
    agree = sum(p == q for p, q in zip(x, y))
    print(f"Items annotated by both: {len(ids)}")
    print(f"Percent agreement (6 types):  {100 * agree / len(ids):.1f}%")
    print(f"Cohen's kappa (6 types):      {k:.3f}  95% CI [{lo:.3f}, {hi:.3f}]")
    print(f"Cohen's kappa (F1 vs rest):   {kb:.3f}  95% CI [{lob:.3f}, {hib:.3f}]\n")
    print_matrix(matrix(x, y, TYPES), TYPES, "A", "B")
    adj = []
    for i in common:
        a_t, b_t = A[i]["failure_type"], B[i]["failure_type"]
        if i in partial:
            how = "ONE_ANNOTATOR_NEEDS_DECISION" if (a_t or b_t) else "UNANNOTATED_NEEDS_DECISION"
        else:
            how = "agreed" if a_t == b_t else "NEEDS_ADJUDICATION"
        src = B[i] if B[i].get("question") else A[i]       # sheet text as sent (one annotator edited a few cells)
        golds = [g for g in (src.get("gold_answers") or "").split(" | ") if g.strip()]
        adj.append({"item_id": i, "annot_A": a_t, "annot_B": b_t, "final_type": a_t if how == "agreed" else "",
                    "how": how, "rationale_A": A[i].get("rationale", ""), "rationale_B": B[i].get("rationale", ""),
                    "question": src.get("question", ""), "gold_answers": src.get("gold_answers", ""),
                    "model_answer": src.get("model_answer", ""), "retrieved_passages": src.get("retrieved_passages", ""),
                    "adjudication_note": "",
                    "model_answer_found_in_passage": found_in(src.get("model_answer", ""), src.get("retrieved_passages", "")),
                    "a_gold_answer_found_in_passage": "yes" if any(found_in(g, src.get("retrieved_passages", "")) == "yes"
                                                                   for g in golds) else "no",
                    **{f"{who}_{short}": S[i].get(col, "") for who, S in (("A", A), ("B", B))
                       for short, col in (("contains", "passage_contains_answer"), ("correct", "answer_correct"),
                                          ("supported", "answer_supported_by_passage"),
                                          ("conflicts", "answer_conflicts_with_passage"))}})
    batch1_cols = ["item_id", "annot_A", "annot_B", "final_type", "how", "rationale_A", "rationale_B", "question",
                   "gold_answers", "model_answer", "retrieved_passages"]
    write_csv(d / "adjudication.csv", adj, batch1_cols)
    (d / "adjudication_rules.txt").write_text(ADJUDICATION_RULES, encoding="utf-8")
    try:
        write_adjudication_xlsx(d / "adjudication.xlsx", adj)
        target = "adjudication.xlsx (yellow cells)"
    except ImportError:
        target = "adjudication.csv"
    json.dump({"n": len(ids), "single_or_unannotated_items": partial, "percent_agreement": agree / len(ids),
               "kappa": k, "kappa_ci": [lo, hi], "kappa_F1_vs_rest": kb, "kappa_F1_vs_rest_ci": [lob, hib],
               "disagreement_matrix": {r: dict(c) for r, c in matrix(x, y, TYPES).items()}},
              open(d / "agreement.json", "w"), indent=2)
    n_dis = sum(r["how"] == "NEEDS_ADJUDICATION" for r in adj)
    print(f"\n{n_dis} disagreements + {len(partial)} single-annotator items to decide -> {target}")
    print("A third person applies the fixed rules (rules tab / adjudication_rules.txt) to every row that is not\n"
          "'agreed', without opening HIDDEN_key.csv. Then run: python Tests/experiment2b_batch2.py report")


# ------------------------------------------------------------------ 3) report
def cmd_report(a):
    d = Path(a.outdir)
    srcs = [p for p in (d / "adjudication.xlsx", d / "adjudication.csv") if p.exists()]
    if not srcs:
        raise SystemExit(f"no adjudication.xlsx / adjudication.csv in {d} -- run the agreement step first")
    src = max(srcs, key=lambda q: q.stat().st_mtime)          # the file the adjudicator saved last
    adj = {r["item_id"]: r for r in read_table(src) if r.get("item_id")}
    print(f"final labels from {src.name}")
    missing = [i for i, r in adj.items() if (r.get("final_type") or "").strip().upper() not in TYPES]
    if missing:
        raise SystemExit(f"{len(missing)} items still lack a valid final_type, e.g. {missing[:5]}")
    for r in adj.values():
        if r.get("how") != "agreed":
            r["how"] = "adjudicated" if r.get("how") in ("NEEDS_ADJUDICATION", "", None) else r["how"]
    key = {r["item_id"]: r for r in read_csv(d / "HIDDEN_key.csv")}
    info = json.load(open(d / "sample_info.json"))
    H = {i: adj[i]["final_type"].strip().upper() for i in adj}
    rep = {}

    print("\nAnnotator agreement by group (before adjudication):")
    for g in ("faithful", "f3", "pilot", "probe"):
        ids = [i for i in adj if key[i]["group"] == g]
        ag = sum(adj[i]["annot_A"] == adj[i]["annot_B"] for i in ids)
        fa = sum(H[i] == adj[i]["annot_A"] for i in ids)
        fb = sum(H[i] == adj[i]["annot_B"] for i in ids)
        print(f"  {g:<9} n={len(ids):<4} A=B {ag:>3} ({ag / max(len(ids), 1):.0%})   final=A {fa:>3}   final=B {fb:>3}")
        rep[f"agreement_{g}"] = {"n": len(ids), "A_equals_B": ag, "final_equals_A": fa, "final_equals_B": fb}
    sug = [i for i in adj if adj[i]["how"] == "adjudicated" and (adj[i].get("claude_suggestion") or "").strip()]
    if sug:
        same = sum(H[i] == adj[i]["claude_suggestion"].strip().upper() for i in sug)
        print(f"Adjudicated items with an AI-drafted suggestion: {len(sug)}; final label = suggestion on {same}, "
              f"changed on {len(sug) - same} (report this in the methods)")
        rep["adjudication_suggestions"] = {"n": len(sug), "accepted": same, "changed": len(sug) - same}
    if "model_answer_found_in_passage" in next(iter(adj.values())):
        print("Rule check (F2 needs the wrong answer to be written in the passage), short answers only:")
        for who in ("A", "B"):
            f2 = [i for i in adj if adj[i][f"annot_{who}"] == "F2" and len(adj[i]["model_answer"].split()) <= 5]
            off = sum(adj[i]["model_answer_found_in_passage"] == "no" for i in f2)
            print(f"  annotator {who}: {off} of {len(f2)} F2 labels are on answers not found in the passage")
            rep[f"rule_check_F2_{who}"] = {"short_answer_F2": len(f2), "answer_not_in_passage": off}

    def dist(ids):
        c = Counter(H[i] for i in ids)
        return ", ".join(f"{t}={c[t]}" for t in TYPES if c[t])

    def ci(k, n):
        lo, hi = wilson(k, n)
        return f"{k}/{n} = {k / n:.2f} [{lo:.2f},{hi:.2f}]" if n else "n/a"

    print("=" * 78 + "\nBATCH 2 RESULTS (human final labels)\n" + "=" * 78)
    print("\nCorrection 6 -- faithful cases (should be F5 = answer correct):")
    for m in MODELS:
        ids = [i for i in adj if key[i]["group"] == "faithful" and key[i]["model"] == m]
        ok = sum(H[i] == "F5" for i in ids)
        print(f"  {m:<6} confirmed faithful {ci(ok, len(ids))}   human types: {dist(ids)}")
        rep[f"faithful_{m}"] = {"n": len(ids), "confirmed": ok, "types": dict(Counter(H[i] for i in ids))}

    print("\nExtra F3 labelling:")
    for m in MODELS:
        for sel in ("predicted_F3", "top_up_p_contains_no"):
            ids = [i for i in adj if key[i]["group"] == "f3" and key[i]["model"] == m and key[i]["selection"] == sel]
            if ids:
                f3 = sum(H[i] == "F3" for i in ids)
                print(f"  {m:<6} {sel:<22} human F3 {ci(f3, len(ids))}   human types: {dist(ids)}")
                rep[f"f3_{m}_{sel}"] = {"n": len(ids), "human_F3": f3, "types": dict(Counter(H[i] for i in ids))}

    print("\nEnrichment pilot -- F1 candidates confirmed by humans (precision):")
    confirmed_rows = []
    for m in MODELS:
        n_pool = info["n_pool_pilot"][m]
        for cond in CONDS:
            ids = [i for i in adj if key[i]["group"] == "pilot" and key[i]["model"] == m and key[i]["cond"] == cond]
            ok = [i for i in ids if H[i] == "F1"]
            probe = [i for i in adj if key[i]["group"] == "probe" and key[i]["model"] == m and key[i]["cond"] == cond]
            missed = sum(float(key[i]["sampling_weight"]) for i in probe if H[i] == "F1")
            per_k = 1000 * len(ok) / n_pool if n_pool else float("nan")
            need = round(40 * n_pool / len(ok)) if ok else None
            print(f"  {m:<6} [{cond}] precision {ci(len(ok), len(ids))}   validated per 1,000 questions {per_k:.1f}"
                  f"   questions for 40 validated: {need}")
            print(f"         probe: {sum(H[i] == 'F1' for i in probe)} human F1 in {len(probe)} other wrong answers "
                  f"-> about {missed:.1f} missed overrides in the {n_pool} pilot questions; human types {dist(probe)}")
            rep[f"pilot_{m}_{cond}"] = {"candidates": len(ids), "confirmed": len(ok), "per_1000": per_k,
                                        "questions_for_40_validated": need, "probe_n": len(probe),
                                        "probe_F1": sum(H[i] == "F1" for i in probe), "estimated_missed": missed}
            for i in ok:
                confirmed_rows.append({"item_id": i, "pid": key[i]["question_id"], "model": m, "cond": cond,
                                       "question": adj[i]["question"], "gold_answers": adj[i]["gold_answers"],
                                       "closed_book_answer": key[i]["closed_book_answer"],
                                       "with_passage_answer": adj[i]["model_answer"],
                                       "same_answer_without_passage": "CHECK"})
    write_csv(d / "pilot_confirmed_F1.csv", confirmed_rows or [{"item_id": ""}],
              list((confirmed_rows or [{"item_id": ""}])[0].keys()))
    print("\n  pilot_confirmed_F1.csv lists the confirmed cases with both answers, for the guide's "
          "no-passage-match check.")

    with open(d / "validated_labels_batch2.jsonl", "w", encoding="utf-8") as f:
        for i in sorted(adj):
            k = key[i]
            f.write(json.dumps({"item_id": i, "group": k["group"], "question_id": k["question_id"], "model": k["model"],
                                "cond": k["cond"], "dataset": k["dataset"], "selection": k["selection"],
                                "human_type": H[i], "auto_type": k["auto_type"], "how": adj[i]["how"],
                                "ai_suggestion": (adj[i].get("claude_suggestion") or "").strip(),
                                "sampling_weight": k["sampling_weight"]}) + "\n")
    json.dump(rep, open(d / "batch2_report.json", "w"), indent=2)
    print(f"\nWrote {d / 'validated_labels_batch2.jsonl'} and {d / 'batch2_report.json'}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("step", choices=["sample", "agreement", "report"])
    ap.add_argument("--dataset", default="data/final/Phase_02/analysis_dataset_relabeled.jsonl")
    ap.add_argument("--retrieval", default="data/final/Phase_02/retrieval_results.jsonl")
    ap.add_argument("--exp1", default="data/final/Phase_05/exp1_failure_types.jsonl")
    ap.add_argument("--validated", default="data/final/Phase_05/exp2/validated_labels.jsonl")
    ap.add_argument("--predictions", default="data/final/Phase_05/judge/rebuilt_predictions_all.csv")
    ap.add_argument("--judge-existing", default="data/final/Phase_05/judge/judge_existing.jsonl")
    ap.add_argument("--pilot-dir", default="data/final/Phase_05/pilot")
    ap.add_argument("--outdir", default="data/final/Phase_05/exp2b")
    ap.add_argument("--n-faithful", type=int, default=30, help="per model")
    ap.add_argument("--n-f3", type=int, default=30, help="per model")
    ap.add_argument("--n-probe", type=int, default=5, help="per model and passage condition")
    ap.add_argument("--single", action="store_true", help="agreement with one annotator only")
    a = ap.parse_args()
    {"sample": cmd_sample, "agreement": cmd_agreement, "report": cmd_report}[a.step](a)


if __name__ == "__main__":
    main()
